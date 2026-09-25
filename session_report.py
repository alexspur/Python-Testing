"""
Build a readable Excel workbook for one GUI session.

Reads the files the GUI already writes into a session folder:
    experiment_log_<ts>.csv   (required)
    gui_log_<ts>.txt          (optional)
    shot_log_<ts>.csv         (optional)
and writes session_report_<ts>.xlsx next to them. The CSV files are never
modified. They stay the raw record. The workbook is a derived, human view.

Usage:
    python session_report.py <session folder or experiment_log csv>
    python session_report.py --all <logs root>     (backfill every session)
"""

import csv
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
SECTION = Font(name=FONT, bold=True, size=11, color="1F3864")
RED = Font(name=FONT, color="C00000", bold=True)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
ERR_FILL = PatternFill("solid", start_color="FDE9E7")
WARN_FILL = PatternFill("solid", start_color="FFF4CE")
OK_FILL = PatternFill("solid", start_color="E8F5E9")

TIME_FMT = "hh:mm:ss.000"

# High-rate telemetry. Everything else is an event.
TELEMETRY = {"OPTA_PSI", "WJ_VOLTAGE"}


# --------------------------------------------------------------------- input
def find_files(target):
    p = Path(target)
    if p.is_file():
        folder, exp = p.parent, p
    else:
        folder = p
        hits = sorted(folder.glob("experiment_log_*.csv"))
        if not hits:
            raise FileNotFoundError(f"no experiment_log_*.csv in {folder}")
        exp = hits[0]
    ts = exp.stem.replace("experiment_log_", "")
    gui = folder / f"gui_log_{ts}.txt"
    shot = folder / f"shot_log_{ts}.csv"
    return folder, ts, exp, gui if gui.exists() else None, shot if shot.exists() else None


def read_csv(path):
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        return list(csv.DictReader(f))


def when(row):
    try:
        return datetime.strptime(row["datetime"], "%Y-%m-%d %H:%M:%S.%f")
    except (KeyError, ValueError):
        return None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def flag(v):
    s = str(v).strip().upper()
    if s in ("1", "TRUE", "YES", "ON"):
        return "ON"
    if s in ("0", "FALSE", "NO", "OFF"):
        return "OFF"
    return s or ""


# -------------------------------------------------------- event descriptions
def describe(r):
    """One plain sentence for an event row. Unknown events fall back to
    listing whatever parameters are present."""
    e, s = r["event_type"], r["source"]
    p1, p2, p3, p4, n = (r.get(k, "") for k in ("param1", "param2", "param3", "param4", "notes"))
    if e == "SESSION_START":
        return f"Session started. Next shot number {p1}. GUI version {p2}."
    if e == "SESSION_END":
        return f"Session ended. Shots fired this session: {p1}."
    if e == "DG535_READBACK":
        return f"Laser DG535 read back. Trigger {p1}. {n}"
    if e in ("INTERLOCK_PASS", "INTERLOCK_FAIL", "INTERLOCK_CHECK"):
        return f"Interlock {p1}: {p2.upper()}. {n}".strip()
    if e == "RELAY_COMMAND":
        conf = "" if p3 in ("", "UNKNOWN") else f", confirmed {p3}"
        return f"Relay {p1} commanded {p2}{conf}. Result: {p4}."
    if e == "RELAY_STATE":
        return f"Relay states ({p1}): {n}"
    if e == "FIRE_BLOCKED":
        return f"FIRE BLOCKED: {p1}. {n}".strip()
    if e == "SHOT":
        return f"SHOT {p1} fired (shot {p2} this session). {n}".strip()
    parts = [x for x in (p1, p2, p3, p4) if x]
    text = " | ".join(parts)
    if n:
        text = f"{text}. {n}" if text else n
    return text


# ------------------------------------------------------------------- styling
def header(ws, cols, widths, row=1):
    for i, (name, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = BODY_FONT


def finish_table(ws, ncols, nrows, hdr_row=1):
    if nrows:
        ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(ncols)}{hdr_row + nrows}"


def empty_note(ws, text, row=2):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, italic=True, color="808080")


# -------------------------------------------------------------------- sheets
def sheet_timeline(wb, gui_lines, events):
    """Human-readable history. Uses the on-screen log when present, since it
    already reads like prose, and falls back to described events."""
    ws = wb.create_sheet("Timeline")
    cols = ["Time", "Device", "Message"]
    header(ws, cols, [14, 16, 110])
    rows = []
    if gui_lines:
        pat = re.compile(r"^\[(\d\d:\d\d:\d\d\.\d{3})\]\s*(?:\[([^\]]+)\])?\s*(.*)$")
        for line in gui_lines:
            m = pat.match(line.rstrip("\n"))
            if not m:
                continue
            t, dev, msg = m.groups()
            rows.append((t, dev or "", msg))
    else:
        for r in events:
            d = when(r)
            rows.append((d.strftime("%H:%M:%S.%f")[:-3] if d else "", r["source"], describe(r)))
    for i, (t, dev, msg) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=t)
        ws.cell(row=i, column=2, value=dev)
        c = ws.cell(row=i, column=3, value=msg)
        c.data_type = "s"          # lines like "=== Auto-connect ===" are text, not formulas
        c.alignment = Alignment(wrap_text=True, vertical="top")
        low = msg.lower()
        fill = None
        if "error" in low or "fault" in low or "blocked" in low or "failed" in low:
            fill = ERR_FILL
        elif "warning" in low or "under range" in low:
            fill = WARN_FILL
        elif "shot" in dev.lower() or "passed" in low:
            fill = OK_FILL
        if fill:
            for col in (1, 2, 3):
                ws.cell(row=i, column=col).fill = fill
    body(ws)
    finish_table(ws, 3, len(rows))
    if not rows:
        empty_note(ws, "No on-screen log for this session.")


def sheet_events(wb, events, t0):
    ws = wb.create_sheet("Events")
    cols = ["Time", "Elapsed (s)", "Event", "Device", "Description"]
    header(ws, cols, [14, 11, 18, 16, 100])
    for i, r in enumerate(events, start=2):
        d = when(r)
        c = ws.cell(row=i, column=1, value=d)
        c.number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r["timestamp_sec"])).number_format = "0.000"
        ws.cell(row=i, column=3, value=r["event_type"])
        ws.cell(row=i, column=4, value=r["source"])
        c = ws.cell(row=i, column=5, value=describe(r))
        c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if r["event_type"] in ("ERROR", "FIRE_BLOCKED", "INTERLOCK_FAIL"):
            for col in range(1, 6):
                ws.cell(row=i, column=col).fill = ERR_FILL
        elif r["event_type"] == "SHOT":
            for col in range(1, 6):
                ws.cell(row=i, column=col).fill = OK_FILL
    body(ws)
    finish_table(ws, 5, len(events))


def sheet_pressure(wb, rows):
    ws = wb.create_sheet("Pressure")
    cols = ["Time", "Elapsed (s)", "Pressure (psi)", "Sensor (V)", "ADC counts", "Status"]
    header(ws, cols, [14, 11, 14, 12, 12, 16])
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r["timestamp_sec"])).number_format = "0.000"
        ws.cell(row=i, column=3, value=num(r["param1"])).number_format = "0.00"
        ws.cell(row=i, column=4, value=num(r["param2"])).number_format = "0.000"
        ws.cell(row=i, column=5, value=num(r["param3"])).number_format = "0"
        st = ws.cell(row=i, column=6, value=r["param4"])
        if r["param4"] and r["param4"].upper() not in ("OK", "GOOD"):
            st.fill = WARN_FILL
    body(ws)
    finish_table(ws, 6, len(rows))
    if not rows:
        empty_note(ws, "No pressure samples this session.")
        return
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = "Marx pressure", "psi", "Elapsed (s)"
    ch.height, ch.width = 9, 22
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=len(rows) + 1))
    ch.legend = None
    ch.x_axis.number_format = "0"
    ch.x_axis.tickLblSkip = max(1, len(rows) // 10)
    ws.add_chart(ch, "H2")


def sheet_hv(wb, name, label, rows):
    ws = wb.create_sheet(name)
    cols = ["Time", "Elapsed (s)", "Measured (kV)", "Current (mA)", "HV", "Fault"]
    header(ws, cols, [14, 11, 14, 13, 8, 8])
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r["timestamp_sec"])).number_format = "0.000"
        ws.cell(row=i, column=3, value=num(r["param1"])).number_format = "0.000"
        ws.cell(row=i, column=4, value=num(r["param2"])).number_format = "0.000"
        hv = ws.cell(row=i, column=5, value=flag(r["param3"]))
        if hv.value == "ON":
            hv.font = RED
        fl = flag(r["param4"])
        f = ws.cell(row=i, column=6, value="YES" if fl == "ON" else ("NO" if fl == "OFF" else fl))
        if f.value == "YES":
            f.fill = ERR_FILL
    body(ws)
    finish_table(ws, 6, len(rows))
    if not rows:
        empty_note(ws, f"No {label} readbacks this session.")
        return
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = f"{label} voltage", "kV", "Elapsed (s)"
    ch.height, ch.width = 9, 22
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=len(rows) + 1))
    ch.legend = None
    ch.x_axis.tickLblSkip = max(1, len(rows) // 10)
    ws.add_chart(ch, "H2")


def sheet_relays(wb, rows):
    ws = wb.create_sheet("Relays")
    cols = ["Time", "Relay", "Channel", "Commanded", "Confirmed", "Result"]
    header(ws, cols, [14, 22, 10, 12, 12, 10])
    rows = [r for r in rows if r["event_type"] == "RELAY_COMMAND"]
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=r["param1"])
        ws.cell(row=i, column=3, value=r["source"].replace("Relay_", ""))
        cmd = ws.cell(row=i, column=4, value=r["param2"])
        if r["param2"].upper() == "ON":
            cmd.font = RED
        ws.cell(row=i, column=5, value=r["param3"] or "UNKNOWN")
        res = ws.cell(row=i, column=6, value=r["param4"])
        if r["param4"].lower() not in ("ok", ""):
            res.fill = ERR_FILL
    body(ws)
    finish_table(ws, 6, len(rows))
    note_row = len(rows) + 3
    c = ws.cell(row=note_row, column=1,
                value="Confirmed = UNKNOWN means the Numato gives no readback. Commanded is what the GUI sent.")
    c.font = Font(name=FONT, italic=True, color="808080")
    if not rows:
        empty_note(ws, "No relay commands this session.")


def sheet_shots(wb, shot_rows):
    ws = wb.create_sheet("Shots")
    if not shot_rows:
        header(ws, ["Shot"], [40])
        empty_note(ws, "No shots fired this session.")
        return
    # Transposed: one column per shot, one row per setting. Easier to read
    # than 80 columns across.
    keys = list(shot_rows[0].keys())
    ws.column_dimensions["A"].width = 32
    c = ws.cell(row=1, column=1, value="Setting")
    c.font, c.fill = HEAD_FONT, HEAD_FILL
    for j, r in enumerate(shot_rows, start=2):
        h = ws.cell(row=1, column=j, value=f"Shot {r.get('shot_number', '')}")
        h.font, h.fill = HEAD_FONT, HEAD_FILL
        h.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = 26
    for i, k in enumerate(keys, start=2):
        ws.cell(row=i, column=1, value=k).font = BOLD
        for j, r in enumerate(shot_rows, start=2):
            v = r.get(k, "")
            n = num(v)
            cell = ws.cell(row=i, column=j, value=n if n is not None and v.strip() != "" else v)
            cell.alignment = Alignment(horizontal="left")
            cell.border = THIN
    ws.freeze_panes = "B2"
    body(ws)


def sheet_summary(wb, ts, events, pressure, wj1, wj2, shots, gui_lines, n_err):
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 60
    ws["A1"] = f"Session {ts}"
    ws["A1"].font = TITLE

    start = next((r for r in events if r["event_type"] == "SESSION_START"), None)
    end = next((r for r in events if r["event_type"] == "SESSION_END"), None)
    dg = [r for r in events if r["event_type"] == "DG535_READBACK"]

    row = 3

    def section(title):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = SECTION
        ws.cell(row=row, column=1).border = THIN
        ws.cell(row=row, column=2).border = THIN
        row += 1

    def line(k, v, fmt=None, font=None):
        nonlocal row
        ws.cell(row=row, column=1, value=k).font = BOLD
        c = ws.cell(row=row, column=2, value=v)
        c.alignment = Alignment(horizontal="left")
        if fmt:
            c.number_format = fmt
        if font:
            c.font = font
        row += 1

    section("Session")
    line("Started", when(start) if start else "", "yyyy-mm-dd hh:mm:ss")
    line("Ended", when(end) if end else "not recorded (GUI may have crashed)",
         "yyyy-mm-dd hh:mm:ss")
    if start and end:
        line("Duration (s)", round(num(end["timestamp_sec"]) - num(start["timestamp_sec"]), 1), "0.0")
    line("GUI version", start["param2"] if start else "")
    line("First shot number available", start["param1"] if start else "")
    line("Shots fired", len(shots))
    line("Errors logged", n_err, font=RED if n_err else None)

    section("Connected devices")
    if gui_lines:
        seen = []
        pat = re.compile(r"^\[[\d:.]+\]\s*\[([^\]]+)\]\s*(.*)$")
        for l in gui_lines:
            m = pat.match(l)
            if not m:
                continue
            dev, msg = m.groups()
            if dev == "AutoConnect":
                m2 = re.match(r"(Rigol\d)_VISA CONNECTED", msg)
                dev = m2.group(1) if m2 else None
            elif not msg.lower().startswith("connected"):
                dev = None
            if dev and dev not in seen:
                seen.append(dev)
        line("Devices", ", ".join(seen) if seen else "none recorded")
    else:
        line("Devices", "no on-screen log available")

    section("Marx pressure (psi)")
    n = len(pressure) + 1
    if pressure:
        line("Minimum", f"=MIN(Pressure!C2:C{n})", "0.00")
        line("Maximum", f"=MAX(Pressure!C2:C{n})", "0.00")
        line("Last reading", f"=INDEX(Pressure!C2:C{n},{len(pressure)})", "0.00")
        line("Samples", f"=COUNT(Pressure!C2:C{n})", "0")
    else:
        line("Samples", 0)

    for name, label, rows in (("WJ1", "WJ1 negative supply", wj1), ("WJ2", "WJ2 positive supply", wj2)):
        section(label)
        m = len(rows) + 1
        if rows:
            line("Peak measured (kV)", f"=MAX({name}!C2:C{m})", "0.000")
            line("Peak current (mA)", f"=MAX({name}!D2:D{m})", "0.000")
            line("Samples with HV ON", f'=COUNTIF({name}!E2:E{m},"ON")', "0")
            line("Samples with fault", f'=COUNTIF({name}!F2:F{m},"YES")', "0")
        else:
            line("Samples", 0)

    section("Laser DG535 timing (last readback)")
    if dg:
        last = dg[-1]
        line("Trigger mode", last["param1"])
        for part in last["notes"].split():
            if "=" in part:
                ch, val = part.split("=", 1)
                line(f"Channel {ch}", val.replace("us", " us"))
    else:
        line("Readback", "none this session")

    section("Sheets in this workbook")
    for s, d in (("Timeline", "On-screen log, color coded. Start here."),
                 ("Events", "Every non-telemetry event, one sentence each."),
                 ("Shots", "Full system snapshot per shot, one column per shot."),
                 ("Pressure", "Every Opta sample with chart."),
                 ("WJ1 / WJ2", "Every HV supply readback with chart."),
                 ("Relays", "Every relay command."),
                 ("Raw log", "The unmodified experiment_log CSV.")):
        line(s, d)
    for r in ws.iter_rows(min_row=3):
        for c in r:
            if c.font.name != FONT:
                c.font = BODY_FONT


def sheet_raw(wb, raw):
    ws = wb.create_sheet("Raw log")
    if not raw:
        return
    cols = list(raw[0].keys())
    header(ws, cols, [12, 24, 18, 14, 16, 16, 12, 12, 60])
    for i, r in enumerate(raw, start=2):
        for j, k in enumerate(cols, start=1):
            c = ws.cell(row=i, column=j, value=r.get(k, ""))
            c.data_type = "s"
    body(ws)
    finish_table(ws, len(cols), len(raw))


# ---------------------------------------------------------------------- main
def build(target):
    folder, ts, exp, gui, shot = find_files(target)
    raw = read_csv(exp)
    gui_lines = gui.read_text(encoding="utf-8", errors="replace").splitlines() if gui else []
    shots = read_csv(shot) if shot else []

    events = [r for r in raw if r["event_type"] not in TELEMETRY]
    pressure = [r for r in raw if r["event_type"] == "OPTA_PSI"]
    wj1 = [r for r in raw if r["event_type"] == "WJ_VOLTAGE" and r["source"] == "WJ1"]
    wj2 = [r for r in raw if r["event_type"] == "WJ_VOLTAGE" and r["source"] == "WJ2"]
    relays = [r for r in raw if r["event_type"].startswith("RELAY")]
    n_err = sum(1 for r in events if r["event_type"] in ("ERROR", "FIRE_BLOCKED"))

    wb = Workbook()
    sheet_summary(wb, ts, events, pressure, wj1, wj2, shots, gui_lines, n_err)
    sheet_timeline(wb, gui_lines, events)
    sheet_events(wb, events, None)
    sheet_shots(wb, shots)
    sheet_pressure(wb, pressure)
    sheet_hv(wb, "WJ1", "WJ1 (negative)", wj1)
    sheet_hv(wb, "WJ2", "WJ2 (positive)", wj2)
    sheet_relays(wb, relays)
    sheet_raw(wb, raw)

    out = folder / f"session_report_{ts}.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--all":
        for exp in sorted(Path(sys.argv[2]).rglob("experiment_log_*.csv")):
            try:
                print("wrote", build(exp))
            except Exception as e:
                print("FAILED", exp, e)
    elif len(sys.argv) == 2:
        print("wrote", build(sys.argv[1]))
    else:
        print(__doc__)
        sys.exit(1)
