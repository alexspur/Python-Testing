"""Session report tests.

Every synthetic session is written through DataLogger and the real shot-log
column list, so the CSV layout under test is the one the GUI writes. No
hardware, no GUI, no real logs folder.
"""

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from utils import session_report as sr
from utils.data_logger import DataLogger
from utils.shot_logger import SHOT_COLUMNS
from utils.shot_snapshot import RIGOL_CHANNEL_SETTING_KEYS, RIGOL_SCOPE_SETTING_KEYS

# The BNC575 CONFIG row, with the same keys the GUI's readback writes.
BNC_FLAT = {"period_s": 0.0001, "system_mode": "SING", "trigger_mode": "DIS"}
for _ch in "ABCD":
    BNC_FLAT.update({f"{_ch}_delay_s": 0.0, f"{_ch}_width_s": 5e-5,
                     f"{_ch}_polarity": "NORM", f"{_ch}_enabled": True})

PREAMBLE = {"format": 0, "points": 1000, "xincrement": 4e-10, "xorigin": -2e-4,
            "yincrement": 0.04, "yorigin": 0.0, "yreference": 128.0}


def scope_settings(n):
    s = {k: "0" for k in RIGOL_SCOPE_SETTING_KEYS}
    s.update({"model": "DS7054", "serial": f"FAKE{n}", "firmware": "00.01",
              "memory_depth": "1.0000E+06", "sample_rate_sa_s": "2.5E+9",
              "timebase_scale_s_div": "5E-7", "timebase_offset_s": "1.83E-4",
              "trigger_source": "CHAN1", "trigger_slope": "POS",
              "trigger_level_v": "0.5", "trigger_sweep": "SING"})
    return s


def channel_settings(n, ch, probe="10"):
    c = {k: "0" for k in RIGOL_CHANNEL_SETTING_KEYS}
    c.update({"display": "1", "scale_v_div": "35", "offset_v": "0", "probe_ratio": probe,
              "coupling": "DC", "impedance": "OMEG", "bandwidth_limit": "OFF",
              "invert": "0", "units": "VOLT", "label": f"CH{ch}"})
    return c


def make_session(tmp_path, shots=(40, 41), gui_log=True, session_end=True,
                 clipped=True, probe_change=True):
    """Write one session. Scope 1's files are on disk and exported OK; scope
    2's export FAILED and its file is absent; scope 3's export is logged OK
    but the file is missing from the folder."""
    dl = DataLogger(log_dir=str(tmp_path / "logs"))
    ts = dl.session_timestamp
    sdir = Path(dl.get_session_dir())
    dl.log_session_start(shots[0] if shots else 40, "abc1234")
    dl.log_connect("BNC575", "COM5", "BNC,575-4")
    dl.log_connect("Rigol1", "TCPIP0::192.168.10.51::5555::SOCKET", "RIGOL,DS7054")
    dl.log_config("BNC575", "all", BNC_FLAT)
    dl.log_config("DG535_laser", "channels", {"A_delay_us": 0.0, "B_delay_us": 180.5,
                                              "trigger_mode": "EXTERNAL"})
    dl.log_config("Opta", "calibration", {"full_scale_psi": 100.0, "zero_offset_mv": 0,
                                          "avg_samples": 32})
    for n in (1, 2, 3):
        dl.log_config(f"Rigol{n}", "scope@connect", scope_settings(n))
        for ch in (1, 2, 3, 4):
            dl.log_config(f"Rigol{n}", f"channel{ch}@connect", channel_settings(n, ch))
    dl.log_relay_state({"charge_positive": True, "charge_negative": False})
    for i in range(5):
        dl.log_opta_pressure(67.5 + i * 0.1, 6.75, 2765)
    dl.log_wj_voltage(1, 70.0, 0.5, hv_on=True)
    dl.log_wj_voltage(2, 70.1, 0.4)

    rows = []
    for idx, shot in enumerate(shots, start=1):
        row = {c: "" for c in SHOT_COLUMNS}
        row.update({"shot_number": str(shot), "session_shot_index": str(idx),
                    "datetime": f"2026-09-25 10:0{idx}:00.000", "pressure_psi": "67.6"})
        for n in (1, 2, 3):
            probe = "20000" if (probe_change and idx == 2 and n == 1) else "10"
            dl.log_config(f"Rigol{n}", "scope@arm", scope_settings(n))
            for k, v in scope_settings(n).items():
                row[f"rigol{n}_{k}"] = v
            for ch in (1, 2, 3, 4):
                cs = channel_settings(n, ch, probe if ch == 1 else "10")
                dl.log_config(f"Rigol{n}", f"channel{ch}@arm", cs)
                for k, v in cs.items():
                    row[f"rigol{n}_ch{ch}_{k}"] = v
            name = f"rigol{n}_{ts}.csv" if idx == 1 else f"rigol{n}_{ts}_shot{idx:02d}.csv"
            row[f"rigol{n}_file"] = name
        rows.append(row)
        dl.log_shot(shot, idx)
        for n in (1, 2, 3):
            dl.log_scope_capture(n, 1000, 1000, shot_number=shot)
            for ch in (1, 2, 3, 4):
                nclip = 5 if (clipped and idx == 1 and n == 1 and ch == 1) else 0
                stats = {"state": "ok", "points": 1000, "clipped": nclip,
                         "clipped_low": nclip, "clipped_high": 0, "code_min": 0 if nclip else 90,
                         "code_max": 200, "v_min": -5.12, "v_max": 5.08, "preamble": PREAMBLE}
                dl.log_scope_channel(n, ch, 1000, stats, shot_number=shot)
                if nclip:
                    dl.log_clip_warning(n, ch, nclip, 1000, shot_number=shot, low=nclip,
                                        code_min=0, code_max=200)
            name = rows[-1][f"rigol{n}_file"]
            if n == 1:
                (sdir / name).write_text("Time (s),Voltage_CH1 (V)\n0,0\n")
                dl.log_scope_export(n, name, 4000, ok=True, shot_number=shot)
            elif n == 2:
                dl.log_scope_export(n, name, 0, ok=False, shot_number=shot, reason="disk full")
                dl.log_error("Rigol2", f"export failed: {name}: disk full")
            else:
                dl.log_scope_export(n, name, 4000, ok=True, shot_number=shot)
    if shots:
        with open(sdir / f"shot_log_{ts}.csv", "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=SHOT_COLUMNS)
            w.writeheader()
            w.writerows(rows)
    if gui_log:
        dl.append_gui_line("[Rigol1] CLIP WARNING: CH1 5 of 1000 samples on the ADC rails")
        dl.append_gui_line("[SHOT] #40 recorded (shot 1 this session)")
    if session_end:
        dl.log_session_end(len(shots))
    return sdir, ts


def sheet_rows(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))


def col(ws, name):
    """1-based column index of a header cell."""
    for c in ws[1]:
        if c.value == name:
            return c.column
    raise KeyError(name)


@pytest.fixture
def session(tmp_path):
    return make_session(tmp_path)


def test_builds_with_two_shots_and_all_sheets(session):
    sdir, ts = session
    out = sr.build(sdir)
    assert out == sdir / f"session_report_{ts}.xlsx" and out.exists()
    wb = load_workbook(out)
    for name in ("Summary", "Timeline", "Events", "Shots", "Scope Settings", "Settings",
                 "Settings History", "Pressure", "WJ1", "WJ2", "Relays", "Raw log"):
        assert name in wb.sheetnames, name
    ws = wb["Shots"]
    assert (ws["B1"].value, ws["C1"].value) == ("Shot 40", "Shot 41")


def test_builds_with_no_shots(tmp_path):
    sdir, _ = make_session(tmp_path, shots=())
    wb = load_workbook(sr.build(sdir))
    assert wb["Shots"]["A2"].value == "No shots fired this session."
    assert sheet_rows(wb["Scope Settings"])[0][0] is None or \
        wb["Scope Settings"]["A2"].value == "No shots this session."


def test_builds_without_a_gui_log_from_events(tmp_path):
    sdir, ts = make_session(tmp_path, gui_log=False)
    wb = load_workbook(sr.build(sdir))
    msgs = [r[2] for r in sheet_rows(wb["Timeline"])]
    assert any("SHOT 40 fired" in str(m) for m in msgs), msgs[:5]


def test_crash_without_session_end_is_reported(tmp_path):
    sdir, _ = make_session(tmp_path, session_end=False)
    wb = load_workbook(sr.build(sdir))
    cells = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert "not recorded" in str(cells["Ended"])


def test_shots_sheet_reads_file_status_from_events_and_disk(session):
    sdir, ts = session
    ws = load_workbook(sr.build(sdir))["Shots"]
    rows = {r[0]: r for r in ws.iter_rows(values_only=True)}
    assert rows["rigol1_export"][1].startswith("OK: 4000 pts")
    assert rows["rigol1_file_on_disk"][1].startswith("on disk")
    assert rows["rigol2_export"][1].startswith("FAILED")
    assert rows["rigol2_file_on_disk"][1] == "MISSING"
    assert rows["rigol3_export"][1].startswith("OK") and rows["rigol3_file_on_disk"][1] == "MISSING"
    assert rows["rigol1_capture"][1] == "CH1 1000 pts, CH2 1000 pts"
    assert rows["rigol1_clipped"][1] == "CH1: 5" and rows["rigol1_clipped"][2] == "0"
    # Red where a file is missing, failed or clipped.
    labels = [c.value for c in ws["A"]]
    for label, shot_col in (("rigol2_file_on_disk", 2), ("rigol2_export", 2),
                            ("rigol3_file_on_disk", 2), ("rigol1_clipped", 2)):
        cell = ws.cell(row=labels.index(label) + 1, column=shot_col)
        assert cell.fill.start_color.rgb.endswith(sr.ERR_FILL.start_color.rgb[-6:]), label
    ok = ws.cell(row=labels.index("rigol1_file_on_disk") + 1, column=2)
    assert ok.fill.fill_type is None
    # The Summary names the missing files.
    summary = {r[0]: r[1] for r in load_workbook(sr.build(sdir))["Summary"].iter_rows(values_only=True) if r[0]}
    assert "rigol2_" in str(summary["Waveform files missing"])
    assert summary["Clipped channel captures"] == 1


def test_scope_settings_has_one_row_per_shot_scope_and_channel(session):
    sdir, ts = session
    ws = load_workbook(sr.build(sdir))["Scope Settings"]
    rows = [r for r in sheet_rows(ws) if r[0] is not None and isinstance(r[0], (int, float))]
    assert len(rows) == 2 * 3 * 4
    assert [(r[0], r[1], r[2]) for r in rows[:5]] == [
        (40, 1, 1), (40, 1, 2), (40, 1, 3), (40, 1, 4), (40, 2, 1)]
    first = dict(zip([c.value for c in ws[1]], rows[0]))
    assert first["Scale (V/div)"] == 35 and first["Probe"] == 10
    assert first["Range min (V)"] == -5.12 and first["Range max (V)"] == 5.08
    assert first["Timebase (s/div)"] == 5e-7 and first["Trigger source"] == "CHAN1"
    assert first["Waveform file"] == f"rigol1_{ts}.csv"
    assert first["File status"].startswith("OK: 4000 pts") and "on disk" in first["File status"]
    assert first["Clipped samples"] == 5


def test_clipped_channel_is_red_and_probe_change_is_yellow(session):
    sdir, _ = session
    ws = load_workbook(sr.build(sdir))["Scope Settings"]
    clip_col, probe_col, status_col = col(ws, "Clipped samples"), col(ws, "Probe"), col(ws, "File status")
    err, warn = sr.ERR_FILL.start_color.rgb[-6:], sr.WARN_FILL.start_color.rgb[-6:]
    # Row 2 is shot 40, scope 1, channel 1: clipped.
    assert ws.cell(row=2, column=clip_col).fill.start_color.rgb.endswith(err)
    assert ws.cell(row=3, column=clip_col).fill.fill_type is None
    # Shot 41, scope 1, channel 1 is row 2 + 12: probe went 10 -> 20000.
    assert ws.cell(row=14, column=probe_col).fill.start_color.rgb.endswith(warn)
    assert ws.cell(row=15, column=probe_col).fill.fill_type is None
    # Scope 2's failed export is red in the status column (row 6 = shot 40, scope 2, ch 1).
    assert ws.cell(row=6, column=status_col).fill.start_color.rgb.endswith(err)


def test_settings_sheet_has_every_bnc_and_rigol_channel_setting(session):
    sdir, _ = session
    ws = load_workbook(sr.build(sdir))["Settings"]
    rows = [dict(zip([c.value for c in ws[1]], r)) for r in sheet_rows(ws)]
    keyed = {(r["Device"], r["Group"], r["Setting"]): r for r in rows}
    for key in BNC_FLAT:
        assert ("BNC575", "all", key) in keyed, key
    assert keyed[("BNC575", "all", "A_width_s")]["Unit"] == "s"
    assert keyed[("BNC575", "all", "A_width_s")]["Value"] == 5e-5
    for n in (1, 2, 3):
        for ch in (1, 2, 3, 4):
            for key in RIGOL_CHANNEL_SETTING_KEYS:
                assert (f"Rigol{n}", f"channel{ch}", key) in keyed, (n, ch, key)
        for key in RIGOL_SCOPE_SETTING_KEYS:
            assert (f"Rigol{n}", "scope", key) in keyed, (n, key)
    r = keyed[("Rigol1", "channel1", "scale_v_div")]
    assert (r["Unit"], r["Source"], r["Read at"]) == ("V/div", "readback", "arm")
    # Latest wins: shot 41 re-armed scope 1 channel 1 with a 20000:1 probe.
    assert keyed[("Rigol1", "channel1", "probe_ratio")]["Value"] == 20000
    assert keyed[("Relay", "state", "charge_positive")]["Value"] == "ON"
    # Device order: BNC575 first, relays last.
    devices = [r["Device"] for r in rows]
    assert devices[0] == "BNC575" and devices[-1] == "Relay"
    assert devices.index("DG535_laser") < devices.index("Opta") < devices.index("Rigol1")


def test_settings_history_lists_every_config_row(session):
    sdir, _ = session
    wb = load_workbook(sr.build(sdir))
    n_config = sum(1 for r in sheet_rows(wb["Raw log"]) if r[2] == "CONFIG")
    assert len(sheet_rows(wb["Settings History"])) == n_config > 0


def test_events_get_plain_sentences():
    r = {"event_type": "SCOPE_EXPORT", "source": "Rigol2", "param1": "2",
         "param2": "rigol2_x.csv", "param3": "0", "param4": "18",
         "notes": "FAILED: rigol2_x.csv not written (disk full)"}
    assert sr.describe(r) == "Rigol 2 waveform file rigol2_x.csv for shot 18: FAILED: rigol2_x.csv not written (disk full)"
    r = {"event_type": "CLIP_WARNING", "source": "Rigol1", "param1": "1", "param2": "3",
         "param3": "12", "param4": "18", "notes": "CH3: 12 of 1000 samples on the ADC rails"}
    assert sr.describe(r).startswith("CLIPPED: Rigol 1 CH3, 12 samples on the ADC rails for shot 18.")
    r = {"event_type": "WJ_VOLTAGE", "source": "WJ1", "param1": "70.000", "param2": "0.500",
         "param3": "1", "param4": "0", "notes": ""}
    assert sr.describe(r) == "WJ1 reads 70.000 kV, 0.500 mA, HV ON, fault NO."
    r = {"event_type": "CONNECT", "source": "Rigol1", "param1": "TCPIP0::x", "param2": "",
         "param3": "", "param4": "", "notes": "RIGOL,DS7054"}
    assert sr.describe(r) == "Rigol1 connected on TCPIP0::x, id RIGOL,DS7054."
    r = {"event_type": "TIMING", "source": "Rigol1", "param1": "", "param2": "", "param3": "",
         "param4": "18", "notes": "capture tid=1 | arm +0.000s"}
    assert sr.describe(r) == "Capture timing for shot 18: capture tid=1 | arm +0.000s"
    r = {"event_type": "OPTA_PSI", "source": "Opta", "param1": "67.60", "param2": "6.760",
         "param3": "2769", "param4": "OK", "notes": ""}
    assert sr.describe(r) == "Dome pressure 67.60 psi (6.760 V, 2769 counts), sensor OK."
    r = {"event_type": "LASER_ARM", "source": "Laser1", "param1": "ARM", "param2": "ARMED",
         "param3": "", "param4": "", "notes": ""}
    assert sr.describe(r) == "Laser1 ARM, state ARMED."


def test_truncated_last_row_does_not_crash(session):
    sdir, ts = session
    exp = sdir / f"experiment_log_{ts}.csv"
    with open(exp, "a", newline="") as f:
        f.write("99.000000,2026-09-25 10:05:00.000,ERROR")        # cut off mid-row
    wb = load_workbook(sr.build(sdir))
    events = [r[2] for r in sheet_rows(wb["Events"])]
    assert events[-1] == "ERROR"


def test_legacy_telemetry_stays_out_of_the_events_sheet(session):
    sdir, ts = session
    exp = sdir / f"experiment_log_{ts}.csv"
    with open(exp, "a", newline="") as f:
        for i in range(20):
            f.write(f"50.{i:06d},2026-09-25 10:04:{i:02d}.000,GLASSMAN_VOLTAGE,WJ1,1.0,0.0,,,\n")
            f.write(f"51.{i:06d},2026-09-25 10:04:{i:02d}.000,MARX_CHARGE,WJ2,1.0,0.0,,,\n")
    wb = load_workbook(sr.build(sdir))
    assert not any(r[2] in ("GLASSMAN_VOLTAGE", "MARX_CHARGE") for r in sheet_rows(wb["Events"]))
    assert sum(1 for r in sheet_rows(wb["Raw log"]) if r[2] == "GLASSMAN_VOLTAGE") == 20


def test_csv_and_text_logs_are_byte_identical_after_the_report(session):
    sdir, _ = session
    files = sorted(p for p in sdir.iterdir() if p.suffix in (".csv", ".txt"))
    before = {p.name: p.read_bytes() for p in files}
    sr.build(sdir)
    sr.build(sdir)                                    # a rebuild changes nothing either
    assert {p.name: p.read_bytes() for p in files} == before


def test_cli_all_backfills_a_tree(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    make_session(tmp_path / "a")
    make_session(tmp_path / "b", shots=())
    assert sr.main(["--all", str(tmp_path)]) == 0
    assert len(list(tmp_path.rglob("session_report_*.xlsx"))) == 2
    assert sr.main([]) == 2
