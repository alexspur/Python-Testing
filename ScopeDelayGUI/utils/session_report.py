# utils/session_report.py
"""
Build a readable Excel workbook for one GUI session.

Reads the files the GUI already writes into a session folder:
    experiment_log_<ts>.csv   (required)
    gui_log_<ts>.txt          (optional)
    shot_log_<ts>.csv         (optional)
and writes session_report_<ts>.xlsx next to them. The CSV files are never
modified. They stay the raw record. The workbook is a derived, human view.

Sheets:
    Summary           session facts, device list, per-supply peaks, sheet guide
    Timeline          the on-screen GUI log, colour coded (events if no GUI log)
    Events            every non-telemetry event, one sentence each
    Shots             the shot log transposed, one column per shot, plus the
                      capture, export, on-disk and clipping status of every
                      waveform file from SCOPE_CAPTURE / SCOPE_EXPORT /
                      CLIP_WARNING events and a check against the folder
    Scope Settings    one row per shot, scope and channel: the arm-time
                      channel, timebase and trigger settings, the capturable
                      voltage range, the waveform file and its status, and
                      the clipped sample count
    Settings          the latest value of every device setting from CONFIG
    Settings History  every CONFIG row
    Pressure / WJ1 / WJ2 / Relays / Raw log

Usage:
    python -m utils.session_report <session folder or experiment_log csv>
    python -m utils.session_report --all <logs root>     (backfill every session)
"""

import csv
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
SECTION = Font(name=FONT, bold=True, size=11, color="1F3864")
RED = Font(name=FONT, color="C00000", bold=True)
NOTE_FONT = Font(name=FONT, italic=True, color="808080")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
ERR_FILL = PatternFill("solid", start_color="FDE9E7")
WARN_FILL = PatternFill("solid", start_color="FFF4CE")
OK_FILL = PatternFill("solid", start_color="E8F5E9")

TIME_FMT = "hh:mm:ss.000"

# High-rate telemetry. Everything else is an event. GLASSMAN_VOLTAGE and
# MARX_CHARGE are the names the 2025 GUI used for the WJ readbacks; without
# them here an old session's Events sheet is thousands of identical rows.
TELEMETRY = {"OPTA_PSI", "WJ_VOLTAGE", "ARDUINO_PSI", "GLASSMAN_VOLTAGE", "MARX_CHARGE"}

# Device order on the Settings sheet. Unknown sources go after these.
DEVICE_ORDER = ["BNC575", "DG535_laser", "DG535", "WJ1", "WJ2", "Laser1", "Laser2",
                "Opta", "Rigol1", "Rigol2", "Rigol3", "Relay"]

# Unit from the setting key's suffix, longest suffix first. Keys that carry
# no unit in their name get a blank, never a guess.
UNIT_SUFFIXES = [
    ("_s_div", "s/div"), ("_v_div", "V/div"), ("_sa_s", "Sa/s"),
    ("_us", "us"), ("_ns", "ns"), ("_ms", "ms"), ("_mv", "mV"), ("_kv", "kV"),
    ("_ma", "mA"), ("_psi", "psi"), ("_hz", "Hz"), ("_pts", "pts"),
    ("_s", "s"), ("_v", "V"),
]

SCOPES = (1, 2, 3)
CHANNELS = (1, 2, 3, 4)


# --------------------------------------------------------------------- input
def find_files(target):
    p = Path(target)
    if p.is_file():
        folder, exp = p.parent, p
    else:
        folder = p
        hits = sorted(folder.glob("experiment_log_*.csv"))
        if not hits:
            raise FileNotFoundError(f"no experiment_log_*.csv in {folder}")
        exp = hits[0]
    ts = exp.stem.replace("experiment_log_", "")
    gui = folder / f"gui_log_{ts}.txt"
    shot = folder / f"shot_log_{ts}.csv"
    return folder, ts, exp, gui if gui.exists() else None, shot if shot.exists() else None


def read_csv(path):
    """Rows as dicts with every value a string.

    A GUI that was killed mid-write leaves a truncated last row; DictReader
    fills its missing fields with None and every sheet then calls string
    methods on it. Such a row is padded with blanks, and a row with no event
    type at all is dropped.
    """
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = []
        for r in csv.DictReader(f):
            r.pop(None, None)                       # fields beyond the header
            rows.append({k: ("" if v is None else v) for k, v in r.items()})
    # A shot-log row (no event_type column) is kept even with a blank shot
    # number: that is what a shot fired under a locked counter looks like.
    return [r for r in rows if "event_type" not in r or r["event_type"] != ""]


def when(row):
    try:
        return datetime.strptime(row["datetime"], "%Y-%m-%d %H:%M:%S.%f")
    except (KeyError, ValueError):
        return None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def flag(v):
    s = str(v).strip().upper()
    if s in ("1", "TRUE", "YES", "ON"):
        return "ON"
    if s in ("0", "FALSE", "NO", "OFF"):
        return "OFF"
    return s or ""


def unit_for(key):
    k = key.lower()
    for suffix, unit in UNIT_SUFFIXES:
        if k.endswith(suffix):
            return unit
    return ""


def parse_settings(notes):
    """'key=value; key=value' (a CONFIG row's notes) -> list of (key, value)."""
    out = []
    for part in notes.split("; "):
        part = part.strip().rstrip(";")
        if "=" in part:
            k, v = part.split("=", 1)
            out.append((k.strip(), v.strip()))
    return out


def _for_shot(p4):
    return f" for shot {p4}" if p4 else ""


# -------------------------------------------------------- event descriptions
def describe(r):
    """One plain sentence for an event row.

    Every event type utils/data_logger.py writes is covered, with each
    parameter's meaning taken from the logger method that writes it. Unknown
    events fall back to listing whatever parameters are present.
    """
    e, s = r.get("event_type", ""), r.get("source", "")
    p1, p2, p3, p4, n = (r.get(k, "") for k in ("param1", "param2", "param3", "param4", "notes"))
    tail = f" {n}" if n else ""

    if e == "SESSION_START":
        return f"Session started. Next shot number {p1}. GUI version {p2}.{tail}"
    if e == "SESSION_END":
        return f"Session ended. Shots fired this session: {p1}.{tail}"
    if e == "SHOT":
        return f"SHOT {p1} fired (shot {p2} this session).{tail}"
    if e == "FIRE_BLOCKED":
        return f"FIRE BLOCKED: {p1}.{tail}"
    if e in ("INTERLOCK_PASS", "INTERLOCK_FAIL", "INTERLOCK_CHECK"):
        step = f" {p1}" if p1 else ""
        verdict = f": {p2.upper()}" if p2 else ""
        return f"Interlock{step}{verdict}.{tail}"
    if e == "ARDUINO_PSI":
        return f"Arduino pressure: CH0 {p1} psi, CH1 {p2} psi, CH2 {p3} psi."
    if e == "ARDUINO_SWITCH":
        return f"Arduino output DO{p1} set {flag(p2) or p2}."
    if e == "WJ_VOLTAGE":
        return (f"{s} reads {p1} kV, {p2} mA, HV {flag(p3)}, fault "
                f"{'YES' if flag(p4) == 'ON' else 'NO'}.{tail}")
    if e == "WJ_COMMAND":
        return f"{s} command {p1} {p2}.".replace("  ", " ")
    if e == "OPTA_PSI":
        return f"Dome pressure {p1} psi ({p2} V, {p3} counts), sensor {p4}."
    if e == "DG535_PULSE":
        return f"DG535 pulse fired: delay {p1} s, width {p2} s."
    if e == "DG535_CONFIG":
        return f"DG535 configured: delay {p1} s, width {p2} s."
    if e == "DG535_READBACK":
        return f"Laser DG535 read back. Trigger {p1}.{tail}"
    if e == "BNC575_PULSE":
        return f"BNC575 fired ({p1}).{tail}"
    if e == "BNC575_ARM":
        return f"BNC575 armed for external trigger at {p1} V."
    if e == "BNC575_CONFIG":
        return f"BNC575 configured: {n or f'A width {p1} s delay {p2} s, B width {p3} s delay {p4} s'}"
    if e == "RELAY_COMMAND":
        conf = "" if p3 in ("", "UNKNOWN") else f", confirmed {p3}"
        return f"Relay {p1} commanded {p2}{conf}. Result: {p4}."
    if e == "RELAY_STATE":
        return f"Relay states ({p1}): {n}"
    if e.startswith("LASER_"):
        state = f", state {p2}" if p2 else ""
        return f"{s} {p1 or e[6:]}{state}.{tail}"
    if e == "SCOPE_CAPTURE":
        return f"Rigol {p1} captured: CH1 {p2} points, CH2 {p3} points{_for_shot(p4)}."
    if e == "SCOPE_EXPORT":
        return f"Rigol {p1} waveform file {p2}{_for_shot(p4)}: {n}"
    if e == "SCOPE_CHANNEL":
        return f"Rigol {p1} CH{p2}: {p3} points{_for_shot(p4)}.{tail}"
    if e == "CLIP_WARNING":
        return (f"CLIPPED: Rigol {p1} CH{p2}, {p3} samples on the ADC rails"
                f"{_for_shot(p4)}.{tail}")
    if e == "SCOPE_ALL":
        return "All armed scopes have finished."
    if e == "SCOPE_ARM":
        return f"Rigol {p1} armed (single trigger)."
    if e == "CONFIG":
        return f"{s} {p1} settings ({p2}): {n}"
    if e == "CONNECT":
        idn = f", id {n}" if n else ""
        return f"{s} connected on {p1}{idn}."
    if e == "DISCONNECT":
        return f"{s} disconnected" + (f": {n}" if n else ".")
    if e == "TIMING":
        return f"Capture timing{_for_shot(p4)}: {n}"
    if e == "ANALYSIS":
        err = f" Error: {n}" if n and "error=" in n and not n.endswith("error=") else ""
        return (f"Analysis of shot {p4}: {p1}. Spacing commanded {p2} ns, "
                f"Q-switch {p3} ns, {n}.{err}").replace("..", ".")
    if e == "ERROR":
        return f"ERROR: {n}" if n else "ERROR"
    if e == "INFO":
        return n
    parts = [x for x in (p1, p2, p3, p4) if x]
    text = " | ".join(parts)
    if n:
        text = f"{text}. {n}" if text else n
    return text


# ------------------------------------------------------------------- styling
def header(ws, cols, widths, row=1):
    for i, (name, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = BODY_FONT


def finish_table(ws, ncols, nrows, hdr_row=1):
    if nrows:
        ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(ncols)}{hdr_row + nrows}"


def empty_note(ws, text, row=2):
    c = ws.cell(row=row, column=1, value=text)
    c.font = NOTE_FONT


def fill_row(ws, row, ncols, fill):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = fill


# ------------------------------------------------- per-shot event digest
class ShotFacts:
    """What the event log says about each shot's waveform files.

    Keyed by shot number as written in param4 of SCOPE_CAPTURE, SCOPE_EXPORT,
    SCOPE_CHANNEL and CLIP_WARNING (the same string the SHOT row carries in
    param1). A shot with no such events has no entry.
    """

    def __init__(self, events, folder):
        self.folder = Path(folder)
        self.capture = {}    # (shot, scope) -> last SCOPE_CAPTURE row
        self.export = {}     # (shot, scope) -> last SCOPE_EXPORT row
        self.channel = {}    # (shot, scope, ch) -> last SCOPE_CHANNEL row
        self.clip = {}       # (shot, scope, ch) -> last CLIP_WARNING row
        for r in events:
            e, shot = r.get("event_type"), r.get("param4", "")
            scope = r.get("param1", "")
            if e == "SCOPE_CAPTURE":
                self.capture[(shot, scope)] = r
            elif e == "SCOPE_EXPORT":
                self.export[(shot, scope)] = r
            elif e == "SCOPE_CHANNEL":
                self.channel[(shot, scope, r.get("param2", ""))] = r
            elif e == "CLIP_WARNING":
                self.clip[(shot, scope, r.get("param2", ""))] = r

    def capture_text(self, shot, scope):
        r = self.capture.get((shot, str(scope)))
        if r is None:
            return "none logged"
        return f"CH1 {r['param2']} pts, CH2 {r['param3']} pts"

    def export_text(self, shot, scope):
        r = self.export.get((shot, str(scope)))
        if r is None:
            return "not exported"
        return r["notes"] or ("OK" if r["param3"] not in ("", "0") else "FAILED")

    def on_disk(self, name):
        """'on disk (12.3 MB)', 'MISSING', or '' when the row names no file."""
        if not name:
            return ""
        p = self.folder / name
        if not p.exists():
            return "MISSING"
        return f"on disk ({p.stat().st_size / 1e6:.1f} MB)"

    def channel_values(self, shot, scope, ch):
        r = self.channel.get((shot, str(scope), str(ch)))
        return dict(parse_settings(r["notes"])) if r else {}

    def clipped(self, shot, scope, ch):
        """Clipped sample count, or None when nothing was logged."""
        v = self.channel_values(shot, scope, ch).get("clipped")
        if v is None:
            r = self.clip.get((shot, str(scope), str(ch)))
            v = r["param3"] if r else None
        n = num(v)
        return None if n is None else int(n)

    def clipped_text(self, shot, scope):
        parts, seen = [], False
        for ch in CHANNELS:
            n = self.clipped(shot, scope, ch)
            if n is None:
                continue
            seen = True
            if n:
                parts.append(f"CH{ch}: {n}")
        if not seen:
            return "not logged"
        return ", ".join(parts) if parts else "0"


def missing_files(shots, facts):
    """(shot number, filename) for every named waveform file not on disk."""
    out = []
    for r in shots:
        for n in SCOPES:
            name = r.get(f"rigol{n}_file", "")
            if name and facts.on_disk(name) == "MISSING":
                out.append((r.get("shot_number", ""), name))
    return out


# -------------------------------------------------------------------- sheets
def sheet_timeline(wb, gui_lines, events):
    """Human-readable history. Uses the on-screen log when present, since it
    already reads like prose, and falls back to described events."""
    ws = wb.create_sheet("Timeline")
    cols = ["Time", "Device", "Message"]
    header(ws, cols, [14, 16, 110])
    rows = []
    if gui_lines:
        pat = re.compile(r"^\[(\d\d:\d\d:\d\d\.\d{3})\]\s*(?:\[([^\]]+)\])?\s*(.*)$")
        for line in gui_lines:
            m = pat.match(line.rstrip("\n"))
            if not m:
                continue
            t, dev, msg = m.groups()
            rows.append((t, dev or "", msg))
    else:
        for r in events:
            d = when(r)
            rows.append((d.strftime("%H:%M:%S.%f")[:-3] if d else "", r["source"], describe(r)))
    for i, (t, dev, msg) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=t)
        ws.cell(row=i, column=2, value=dev)
        c = ws.cell(row=i, column=3, value=msg)
        c.data_type = "s"          # lines like "=== Auto-connect ===" are text, not formulas
        c.alignment = Alignment(wrap_text=True, vertical="top")
        low = msg.lower()
        fill = None
        if ("error" in low or "fault" in low or "blocked" in low or "failed" in low
                or "clip" in low):
            fill = ERR_FILL
        elif "warning" in low or "under range" in low:
            fill = WARN_FILL
        elif "shot" in dev.lower() or "passed" in low:
            fill = OK_FILL
        if fill:
            fill_row(ws, i, 3, fill)
    body(ws)
    finish_table(ws, 3, len(rows))
    if not rows:
        empty_note(ws, "No on-screen log for this session.")


def sheet_events(wb, events):
    ws = wb.create_sheet("Events")
    cols = ["Time", "Elapsed (s)", "Event", "Device", "Description"]
    header(ws, cols, [14, 11, 18, 16, 100])
    for i, r in enumerate(events, start=2):
        d = when(r)
        c = ws.cell(row=i, column=1, value=d)
        c.number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r.get("timestamp_sec"))).number_format = "0.000"
        ws.cell(row=i, column=3, value=r["event_type"])
        ws.cell(row=i, column=4, value=r["source"])
        c = ws.cell(row=i, column=5, value=describe(r))
        c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if r["event_type"] in ("ERROR", "FIRE_BLOCKED", "INTERLOCK_FAIL", "CLIP_WARNING"):
            fill_row(ws, i, 5, ERR_FILL)
        elif r["event_type"] == "SHOT":
            fill_row(ws, i, 5, OK_FILL)
    body(ws)
    finish_table(ws, 5, len(events))


def sheet_pressure(wb, rows):
    ws = wb.create_sheet("Pressure")
    cols = ["Time", "Elapsed (s)", "Pressure (psi)", "Sensor (V)", "ADC counts", "Status"]
    header(ws, cols, [14, 11, 14, 12, 12, 16])
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r["timestamp_sec"])).number_format = "0.000"
        ws.cell(row=i, column=3, value=num(r["param1"])).number_format = "0.00"
        ws.cell(row=i, column=4, value=num(r["param2"])).number_format = "0.000"
        ws.cell(row=i, column=5, value=num(r["param3"])).number_format = "0"
        st = ws.cell(row=i, column=6, value=r["param4"])
        if r["param4"] and r["param4"].upper() not in ("OK", "GOOD"):
            st.fill = WARN_FILL
    body(ws)
    finish_table(ws, 6, len(rows))
    if not rows:
        empty_note(ws, "No pressure samples this session.")
        return
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = "Marx pressure", "psi", "Elapsed (s)"
    ch.height, ch.width = 9, 22
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=len(rows) + 1))
    ch.legend = None
    ch.x_axis.number_format = "0"
    ch.x_axis.tickLblSkip = max(1, len(rows) // 10)
    ws.add_chart(ch, "H2")


def sheet_hv(wb, name, label, rows):
    ws = wb.create_sheet(name)
    cols = ["Time", "Elapsed (s)", "Measured (kV)", "Current (mA)", "HV", "Fault"]
    header(ws, cols, [14, 11, 14, 13, 8, 8])
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r["timestamp_sec"])).number_format = "0.000"
        ws.cell(row=i, column=3, value=num(r["param1"])).number_format = "0.000"
        ws.cell(row=i, column=4, value=num(r["param2"])).number_format = "0.000"
        hv = ws.cell(row=i, column=5, value=flag(r["param3"]))
        if hv.value == "ON":
            hv.font = RED
        fl = flag(r["param4"])
        f = ws.cell(row=i, column=6, value="YES" if fl == "ON" else ("NO" if fl == "OFF" else fl))
        if f.value == "YES":
            f.fill = ERR_FILL
    body(ws)
    finish_table(ws, 6, len(rows))
    if not rows:
        empty_note(ws, f"No {label} readbacks this session.")
        return
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = f"{label} voltage", "kV", "Elapsed (s)"
    ch.height, ch.width = 9, 22
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=len(rows) + 1))
    ch.legend = None
    ch.x_axis.tickLblSkip = max(1, len(rows) // 10)
    ws.add_chart(ch, "H2")


def sheet_relays(wb, rows):
    ws = wb.create_sheet("Relays")
    cols = ["Time", "Relay", "Channel", "Commanded", "Confirmed", "Result"]
    header(ws, cols, [14, 22, 10, 12, 12, 10])
    rows = [r for r in rows if r["event_type"] == "RELAY_COMMAND"]
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=r["param1"])
        ws.cell(row=i, column=3, value=r["source"].replace("Relay_", ""))
        cmd = ws.cell(row=i, column=4, value=r["param2"])
        if r["param2"].upper() == "ON":
            cmd.font = RED
        ws.cell(row=i, column=5, value=r["param3"] or "UNKNOWN")
        res = ws.cell(row=i, column=6, value=r["param4"])
        if r["param4"].lower() not in ("ok", ""):
            res.fill = ERR_FILL
    body(ws)
    finish_table(ws, 6, len(rows))
    note_row = len(rows) + 3
    c = ws.cell(row=note_row, column=1,
                value="Confirmed = UNKNOWN means the Numato gives no readback. Commanded is what the GUI sent.")
    c.font = NOTE_FONT
    if not rows:
        empty_note(ws, "No relay commands this session.")


def shot_status_rows(shots, facts):
    """Derived rows for the Shots sheet, keyed by the shot-log column they
    follow. Each value is (text, fill or None), one per shot."""
    derived = {}
    for n in SCOPES:
        rows = []
        for label, fn in (
            (f"rigol{n}_capture", lambda r, n=n: facts.capture_text(r.get("shot_number", ""), n)),
            (f"rigol{n}_export", lambda r, n=n: facts.export_text(r.get("shot_number", ""), n)),
            (f"rigol{n}_file_on_disk", lambda r, n=n: facts.on_disk(r.get(f"rigol{n}_file", ""))),
            (f"rigol{n}_clipped", lambda r, n=n: facts.clipped_text(r.get("shot_number", ""), n)),
        ):
            cells = []
            for r in shots:
                text = fn(r)
                bad = (text == "MISSING" or text.startswith("FAILED")
                       or (label.endswith("_clipped") and text not in ("0", "not logged")))
                cells.append((text, ERR_FILL if bad else None))
            rows.append((label, cells))
        derived[f"rigol{n}_file"] = rows
    return derived


def sheet_shots(wb, shot_rows, facts):
    ws = wb.create_sheet("Shots")
    if not shot_rows:
        header(ws, ["Shot"], [40])
        empty_note(ws, "No shots fired this session.")
        return
    # Transposed: one column per shot, one row per setting. Easier to read
    # than 270 columns across. The derived file-status rows sit under each
    # scope's rigol<N>_file row.
    keys = list(shot_rows[0].keys())
    derived = shot_status_rows(shot_rows, facts)
    lines = []
    for k in keys:
        lines.append((k, [(r.get(k, ""), None) for r in shot_rows], False))
        for label, cells in derived.pop(k, []):
            lines.append((label, cells, True))
    for rows in derived.values():                     # no rigol<N>_file column
        for label, cells in rows:
            lines.append((label, cells, True))

    ws.column_dimensions["A"].width = 32
    c = ws.cell(row=1, column=1, value="Setting")
    c.font, c.fill = HEAD_FONT, HEAD_FILL
    for j, r in enumerate(shot_rows, start=2):
        h = ws.cell(row=1, column=j, value=f"Shot {r.get('shot_number', '')}")
        h.font, h.fill = HEAD_FONT, HEAD_FILL
        h.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = 26
    for i, (label, cells, is_derived) in enumerate(lines, start=2):
        lc = ws.cell(row=i, column=1, value=label)
        lc.font = BOLD
        if is_derived:
            lc.font = Font(name=FONT, bold=True, color="1F3864")
        for j, (v, fill) in enumerate(cells, start=2):
            # Shot-log values become numbers where they parse; the derived
            # status rows are text ("0", "CH1: 5", "MISSING") and stay so.
            n = None if is_derived else num(v)
            cell = ws.cell(row=i, column=j,
                           value=n if n is not None and str(v).strip() != "" else v)
            cell.alignment = Alignment(horizontal="left")
            cell.border = THIN
            if fill is not None:
                cell.fill = fill
    ws.freeze_panes = "B2"
    body(ws)


SCOPE_SETTING_COLS = [
    ("Shot", 8), ("Scope", 7), ("Channel", 8), ("Display", 8),
    ("Scale (V/div)", 13), ("Offset (V)", 12), ("Probe", 8), ("Coupling", 9),
    ("Impedance", 10), ("BW limit", 9), ("Invert", 7), ("Units", 7), ("Label", 10),
    ("Range min (V)", 13), ("Range max (V)", 13),
    ("Timebase (s/div)", 15), ("Timebase offset (s)", 17), ("Sample rate (Sa/s)", 16),
    ("Memory depth", 13), ("Trigger source", 13), ("Trigger slope", 12),
    ("Trigger level (V)", 15), ("Trigger sweep", 12),
    ("Waveform file", 34), ("File status", 30), ("Clipped samples", 14),
]
CHANNEL_KEYS = ["display", "scale_v_div", "offset_v", "probe_ratio", "coupling",
                "impedance", "bandwidth_limit", "invert", "units", "label"]
SCOPE_KEYS = ["timebase_scale_s_div", "timebase_offset_s", "sample_rate_sa_s",
              "memory_depth", "trigger_source", "trigger_slope", "trigger_level_v",
              "trigger_sweep"]


def scope_setting_rows(shots, facts):
    """One dict per shot, scope and channel for the Scope Settings sheet.
    'changed' names the cells whose probe ratio or scale differs from the
    previous shot on the same scope and channel."""
    out, previous = [], {}
    for r in shots:
        shot = r.get("shot_number", "")
        for n in SCOPES:
            name = r.get(f"rigol{n}_file", "")
            status = facts.export_text(shot, n)
            disk = facts.on_disk(name)
            if disk:
                status = f"{status}; {disk}"
            for ch in CHANNELS:
                g = lambda k: r.get(f"rigol{n}_ch{ch}_{k}", "")  # noqa: E731
                cv = facts.channel_values(shot, n, ch)
                row = {"Shot": shot, "Scope": n, "Channel": ch}
                for col, key in zip(["Display", "Scale (V/div)", "Offset (V)", "Probe",
                                     "Coupling", "Impedance", "BW limit", "Invert",
                                     "Units", "Label"], CHANNEL_KEYS):
                    row[col] = g(key)
                row["Range min (V)"] = cv.get("v_min", "")
                row["Range max (V)"] = cv.get("v_max", "")
                for col, key in zip(["Timebase (s/div)", "Timebase offset (s)",
                                     "Sample rate (Sa/s)", "Memory depth", "Trigger source",
                                     "Trigger slope", "Trigger level (V)", "Trigger sweep"],
                                    SCOPE_KEYS):
                    row[col] = r.get(f"rigol{n}_{key}", "")
                row["Waveform file"] = name
                row["File status"] = status
                clipped = facts.clipped(shot, n, ch)
                row["Clipped samples"] = "" if clipped is None else clipped
                changed = set()
                prev = previous.get((n, ch))
                if prev is not None:
                    for col in ("Probe", "Scale (V/div)"):
                        if prev[col] != row[col]:
                            changed.add(col)
                previous[(n, ch)] = row
                row["changed"] = changed
                row["clipped_flag"] = bool(clipped)
                row["missing_flag"] = disk == "MISSING" or status.startswith("FAILED")
                out.append(row)
    return out


def sheet_scope_settings(wb, shots, facts):
    ws = wb.create_sheet("Scope Settings")
    cols = [c for c, _ in SCOPE_SETTING_COLS]
    header(ws, cols, [w for _, w in SCOPE_SETTING_COLS])
    rows = scope_setting_rows(shots, facts)
    numeric = {"Shot", "Scope", "Channel", "Scale (V/div)", "Offset (V)", "Probe",
               "Range min (V)", "Range max (V)", "Timebase (s/div)", "Timebase offset (s)",
               "Sample rate (Sa/s)", "Trigger level (V)", "Clipped samples"}
    for i, row in enumerate(rows, start=2):
        for j, col in enumerate(cols, start=1):
            v = row.get(col, "")
            n = num(v) if col in numeric else None
            cell = ws.cell(row=i, column=j, value=n if n is not None and str(v) != "" else v)
            if col in row["changed"]:
                cell.fill = WARN_FILL
        if row["clipped_flag"]:
            for col in ("Channel", "Clipped samples"):
                ws.cell(row=i, column=cols.index(col) + 1).fill = ERR_FILL
        if row["missing_flag"]:
            ws.cell(row=i, column=cols.index("File status") + 1).fill = ERR_FILL
    body(ws)
    finish_table(ws, len(cols), len(rows))
    note = len(rows) + 3
    ws.cell(row=note, column=1, value=(
        "Red: clipped samples on the ADC rails, or a waveform file that failed or is "
        "missing. Yellow: probe ratio or scale changed since the previous shot on that "
        "channel. Range is the capturable voltage span from the waveform preamble, "
        "logged per capture; blank when no capture was logged.")).font = NOTE_FONT
    if not rows:
        empty_note(ws, "No shots this session.")


def latest_settings(events):
    """Latest value of every (device, group, key) from CONFIG rows plus the
    latest RELAY_STATE. Returns rows sorted for the Settings sheet."""
    latest = {}
    for r in events:
        e = r.get("event_type")
        if e == "CONFIG":
            group, read_at = r["param1"], ""
            if "@" in group:
                group, read_at = group.split("@", 1)
            for key, val in parse_settings(r["notes"]):
                latest[(r["source"], group, key)] = (val, r["param2"], read_at, r)
        elif e == "RELAY_STATE":
            for part in r["notes"].split(", "):
                if "=" in part:
                    name, state = part.split("=", 1)
                    latest[("Relay", "state", name.strip())] = (state.strip(), r["param1"], "", r)

    def order(item):
        (device, group, key), _ = item
        rank = DEVICE_ORDER.index(device) if device in DEVICE_ORDER else len(DEVICE_ORDER)
        return (rank, device, group, key)

    out = []
    for (device, group, key), (val, origin, read_at, r) in sorted(latest.items(), key=order):
        out.append({"Device": device, "Group": group, "Setting": key, "Value": val,
                    "Unit": unit_for(key), "Source": origin, "Read at": read_at,
                    "Time": when(r)})
    return out


def sheet_settings(wb, events):
    ws = wb.create_sheet("Settings")
    cols = ["Device", "Group", "Setting", "Value", "Unit", "Source", "Read at", "Time"]
    header(ws, cols, [12, 14, 24, 40, 8, 11, 10, 14])
    rows = latest_settings(events)
    for i, r in enumerate(rows, start=2):
        for j, col in enumerate(cols, start=1):
            v = r[col]
            if col == "Value":
                n = num(v)
                v = n if n is not None and str(v).strip() != "" else v
            c = ws.cell(row=i, column=j, value=v)
            if col == "Time":
                c.number_format = TIME_FMT
            if col == "Value" and str(r["Value"]).upper() == "UNKNOWN":
                c.fill = WARN_FILL
    body(ws)
    finish_table(ws, len(cols), len(rows))
    if not rows:
        empty_note(ws, "No CONFIG rows this session (GUI older than the CONFIG event).")


def sheet_settings_history(wb, events):
    ws = wb.create_sheet("Settings History")
    cols = ["Time", "Elapsed (s)", "Device", "Group", "Source", "Settings"]
    header(ws, cols, [14, 11, 12, 18, 11, 120])
    rows = [r for r in events if r.get("event_type") == "CONFIG"]
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=when(r)).number_format = TIME_FMT
        ws.cell(row=i, column=2, value=num(r.get("timestamp_sec"))).number_format = "0.000"
        ws.cell(row=i, column=3, value=r["source"])
        ws.cell(row=i, column=4, value=r["param1"])
        ws.cell(row=i, column=5, value=r["param2"])
        c = ws.cell(row=i, column=6, value=r["notes"])
        c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
    body(ws)
    finish_table(ws, 6, len(rows))
    if not rows:
        empty_note(ws, "No CONFIG rows this session.")


def connected_devices(events, gui_lines):
    seen = []
    for r in events:
        if r.get("event_type") == "CONNECT" and r["source"] not in seen:
            seen.append(r["source"])
    if seen:
        return seen
    pat = re.compile(r"^\[[\d:.]+\]\s*\[([^\]]+)\]\s*(.*)$")
    for line in gui_lines:
        m = pat.match(line)
        if not m:
            continue
        dev, msg = m.groups()
        if dev == "AutoConnect":
            m2 = re.match(r"(Rigol\d)_VISA CONNECTED", msg)
            dev = m2.group(1) if m2 else None
        elif not msg.lower().startswith("connected"):
            dev = None
        if dev and dev not in seen:
            seen.append(dev)
    return seen


def sheet_summary(wb, ts, events, pressure, wj1, wj2, shots, gui_lines, n_err, missing,
                  n_clipped):
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 70
    ws["A1"] = f"Session {ts}"
    ws["A1"].font = TITLE

    start = next((r for r in events if r["event_type"] == "SESSION_START"), None)
    end = next((r for r in events if r["event_type"] == "SESSION_END"), None)
    dg = [r for r in events if r["event_type"] == "DG535_READBACK"]
    analyses = [r for r in events if r["event_type"] == "ANALYSIS"]

    row = 3

    def section(title):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = SECTION
        ws.cell(row=row, column=1).border = THIN
        ws.cell(row=row, column=2).border = THIN
        row += 1

    def line(k, v, fmt=None, font=None):
        nonlocal row
        ws.cell(row=row, column=1, value=k).font = BOLD
        c = ws.cell(row=row, column=2, value=v)
        c.alignment = Alignment(horizontal="left")
        if fmt:
            c.number_format = fmt
        if font:
            c.font = font
        row += 1

    section("Session")
    line("Started", when(start) if start else "", "yyyy-mm-dd hh:mm:ss")
    line("Ended", when(end) if end else "not recorded (GUI may have crashed)",
         "yyyy-mm-dd hh:mm:ss")
    if start and end and num(end["timestamp_sec"]) is not None and num(start["timestamp_sec"]) is not None:
        line("Duration (s)", round(num(end["timestamp_sec"]) - num(start["timestamp_sec"]), 1), "0.0")
    line("GUI version", start["param2"] if start else "")
    line("First shot number available", start["param1"] if start else "")
    line("Shots fired", len(shots))
    line("Errors logged", n_err, font=RED if n_err else None)
    line("Waveform files missing",
         (", ".join(f"shot {s}: {f}" for s, f in missing) if missing else 0),
         font=RED if missing else None)
    line("Clipped channel captures", n_clipped, font=RED if n_clipped else None)

    section("Connected devices")
    devs = connected_devices(events, gui_lines)
    line("Devices", ", ".join(devs) if devs else "none recorded")

    section("Marx pressure (psi)")
    n = len(pressure) + 1
    if pressure:
        line("Minimum", f"=MIN(Pressure!C2:C{n})", "0.00")
        line("Maximum", f"=MAX(Pressure!C2:C{n})", "0.00")
        line("Last reading", f"=INDEX(Pressure!C2:C{n},{len(pressure)})", "0.00")
        line("Samples", f"=COUNT(Pressure!C2:C{n})", "0")
    else:
        line("Samples", 0)

    for name, label, rows in (("WJ1", "WJ1 negative supply", wj1), ("WJ2", "WJ2 positive supply", wj2)):
        section(label)
        m = len(rows) + 1
        if rows:
            line("Peak measured (kV)", f"=MAX({name}!C2:C{m})", "0.000")
            line("Peak current (mA)", f"=MAX({name}!D2:D{m})", "0.000")
            line("Samples with HV ON", f'=COUNTIF({name}!E2:E{m},"ON")', "0")
            line("Samples with fault", f'=COUNTIF({name}!F2:F{m},"YES")', "0")
        else:
            line("Samples", 0)

    section("Laser DG535 timing (last readback)")
    if dg:
        last = dg[-1]
        line("Trigger mode", last["param1"])
        for part in last["notes"].split():
            if "=" in part:
                ch, val = part.split("=", 1)
                line(f"Channel {ch}", val.replace("us", " us"))
    else:
        line("Readback", "none this session")

    if analyses:
        section("Shot analysis (last result per shot)")
        last = {}
        for r in analyses:
            last[r["param4"]] = r
        for shot, r in last.items():
            line(f"Shot {shot}", describe(r))

    section("Sheets in this workbook")
    for s, d in (("Timeline", "On-screen log, color coded. Start here."),
                 ("Events", "Every non-telemetry event, one sentence each."),
                 ("Shots", "Full system snapshot per shot, one column per shot, with "
                           "each waveform file's capture, export, on-disk and clipping status."),
                 ("Scope Settings", "One row per shot, scope and channel: arm-time settings, "
                                    "capturable range, file, clipped count."),
                 ("Settings", "Latest value of every device setting (CONFIG rows)."),
                 ("Settings History", "Every CONFIG row in order."),
                 ("Pressure", "Every Opta sample with chart."),
                 ("WJ1 / WJ2", "Every HV supply readback with chart."),
                 ("Relays", "Every relay command."),
                 ("Raw log", "The unmodified experiment_log CSV.")):
        line(s, d)
    for r in ws.iter_rows(min_row=3):
        for c in r:
            if c.font is None or c.font.name != FONT:
                c.font = BODY_FONT


def sheet_raw(wb, raw):
    ws = wb.create_sheet("Raw log")
    if not raw:
        return
    cols = list(raw[0].keys())
    header(ws, cols, [12, 24, 18, 14, 16, 16, 12, 12, 60])
    for i, r in enumerate(raw, start=2):
        for j, k in enumerate(cols, start=1):
            c = ws.cell(row=i, column=j, value=r.get(k, ""))
            c.data_type = "s"
    body(ws)
    finish_table(ws, len(cols), len(raw))


# ---------------------------------------------------------------------- main
def build(target):
    """Build session_report_<ts>.xlsx for one session. Returns its path.
    Reads only; the CSV and text logs are never opened for writing."""
    folder, ts, exp, gui, shot = find_files(target)
    raw = read_csv(exp)
    gui_lines = gui.read_text(encoding="utf-8", errors="replace").splitlines() if gui else []
    shots = read_csv(shot) if shot else []

    events = [r for r in raw if r["event_type"] not in TELEMETRY]
    pressure = [r for r in raw if r["event_type"] == "OPTA_PSI"]
    wj1 = [r for r in raw if r["event_type"] == "WJ_VOLTAGE" and r["source"] == "WJ1"]
    wj2 = [r for r in raw if r["event_type"] == "WJ_VOLTAGE" and r["source"] == "WJ2"]
    relays = [r for r in raw if r["event_type"].startswith("RELAY")]
    n_err = sum(1 for r in events if r["event_type"] in ("ERROR", "FIRE_BLOCKED"))
    n_clipped = sum(1 for r in events if r["event_type"] == "CLIP_WARNING")
    facts = ShotFacts(events, folder)
    missing = missing_files(shots, facts)

    wb = Workbook()
    sheet_summary(wb, ts, events, pressure, wj1, wj2, shots, gui_lines, n_err, missing,
                  n_clipped)
    sheet_timeline(wb, gui_lines, events)
    sheet_events(wb, events)
    sheet_shots(wb, shots, facts)
    sheet_scope_settings(wb, shots, facts)
    sheet_settings(wb, events)
    sheet_settings_history(wb, events)
    sheet_pressure(wb, pressure)
    sheet_hv(wb, "WJ1", "WJ1 (negative)", wj1)
    sheet_hv(wb, "WJ2", "WJ2 (positive)", wj2)
    sheet_relays(wb, relays)
    sheet_raw(wb, raw)

    out = folder / f"session_report_{ts}.xlsx"
    wb.save(out)
    return out


def main(argv=None):
    argv = sys.argv[1:] if argv is None else list(argv)
    if len(argv) >= 2 and argv[0] == "--all":
        built = failed = 0
        for exp in sorted(Path(argv[1]).rglob("experiment_log_*.csv")):
            try:
                print("wrote", build(exp))
                built += 1
            except Exception as e:  # noqa: BLE001
                print("FAILED", exp, f"{type(e).__name__}: {e}")
                failed += 1
        print(f"{built} built, {failed} failed")
        return 1 if failed else 0
    if len(argv) == 1:
        print("wrote", build(argv[0]))
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main())
